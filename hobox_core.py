# -*- coding: utf-8 -*-
"""호칭박스 핵심 로직 (화면 없이도 쓸 수 있는 함수 모음)

기사 전문에서 지정발언자 한 사람의 호칭별 기간을 찾는다.
- 근거는 [이름+호칭]이 말하는 사람으로 쓰이고 같은 문장에 큰따옴표 인용이 있는 경우만 쓴다.
- 인용문 안의 [이름+호칭], '당시 ○○', 앞뒤와 60일 넘게 떨어진 외톨이 1건은 쓰지 않는다.
- 1단계: 예시 호칭 리스트에 맞는 호칭(현직/전)별 시작일·종료일. 2단계: 예시에 없는 호칭 후보.
- 3단계: 사용자가 2단계에서 고른 호칭을 예시에 더해 다시 계산(analyze의 adopted).
"""
import csv, re, collections, threading
from datetime import datetime, timedelta
import io

JOSA = ['에게서', '에게', '께서', '으로', '로서', '로써', '이라고', '이라', '라고', '이자', '이며', '이나', '까지', '부터',
        '은', '는', '이', '가', '도', '의', '을', '를', '과', '와', '께', '만', '에', '로', '씨', '라']
TITLE_SUFFIX = re.compile(r'(위원장|위원|의원|대변인|대표|후보자|후보|당선인|당선자|장관|차관|총리|대통령|시장|지사|군수|구청장|교육감|총장|의장|원장|이사장|실장|본부장|수석|비서관|보좌관|고문|교수|변호사|작가|아나운서|앵커|대행|단장|간사|특보|사무총장|처장|청장|국장|회장|팀장|센터장|판사|검사|기자|감독|선수|가수|배우)$')
RETRO = re.compile(r'당시|였던|지낸|출신|전직|역임|맡았던|했던')

SENT_END = re.compile(r'(다|요|죠|까|니다|라고|했다|한다)[.”"\'’)\]]*$|[.?!][”"\'’)\]]*$')

def body_without_headings(text):
    """편집인 기준: 줄(문단) 바꿈으로 구분된 덩어리가 문장 종결로 끝나지 않으면 소제목·서명·사진설명이다.
    같은 줄 안에서 띄어쓰기만으로 이어지는 문장(주어 생략 문장 등)은 한 덩어리로 본다.
    기사 맨 끝의 긴 덩어리는 본문이 잘린 것일 수 있어 남긴다."""
    lines = [ln.strip() for ln in text.split('\n') if ln.strip()]
    keep = []
    for i, p in enumerate(lines):
        if SENT_END.search(p):
            keep.append(p)
        elif i == len(lines) - 1 and len(p) > 80:
            keep.append(p)
    return '\n'.join(keep)

def norm_date(s):
    s = re.sub(r'\D', '', s)[:8]
    try:
        return datetime.strptime(s, '%Y%m%d').date()
    except ValueError:
        return None

def strip_josa(tok):
    changed = True
    while changed and len(tok) > 1:
        changed = False
        for j in JOSA:
            if tok.endswith(j) and len(tok) - len(j) >= 2:
                cand = tok[:-len(j)]
                if TITLE_SUFFIX.search(cand) or j in ('에게서', '에게', '께서', '으로', '로서', '로써'):
                    tok = cand
                    changed = True
                    break
    return tok

GAP = 60   # 같은 호칭이 이 일수 넘게 안 나오면 다른 구간으로 본다
TOK = r'[가-힣0-9]{1,12}'

def candidates(text, name):
    """(전여부, 호칭, 회고여부, 위치) 목록. 풀네임 바로 뒤의 낱말."""
    out = []
    for m in re.finditer(re.escape(name) + r'(?![가-힣])\s?(?:\([^)]{0,12}\)\s?)?((?:전|前|현)\s?)?(' + TOK + r')(?:\s(' + TOK + r'))?', text):
        pre = (m.group(1) or '').strip()
        t1 = strip_josa(m.group(2))
        t2 = m.group(3)
        title = None
        lastraw = None
        if TITLE_SUFFIX.search(t1) and len(t1) >= 2:
            title = t1
            lastraw = m.group(2)
            if t2 and strip_josa(t2) in ('후보', '후보자', '예비후보', '내정자', '지명자'):
                title = t1 + ' ' + strip_josa(t2)          # '최고위원 후보'는 최고위원이 아니다
                lastraw = t2
        elif t2:                                        # 앞 낱말이 기관명일 수 있음(청와대 대변인)
            t2s = strip_josa(t2)
            if TITLE_SUFFIX.search(t2s) and len(t2s) >= 2 and not m.group(2).endswith(tuple('은는이가도의을를')):
                title = t1 + ' ' + t2s
                lastraw = t2
        if title:
            lst = title.split()[-1]
            josa = lastraw[len(lst):] if lastraw.startswith(lst) else ''
            end = m.start(3) + len(lst) if (lastraw is t2 or lastraw == t2) and m.group(3) else m.start(2) + len(lst)
            ctx = text[max(0, m.start() - 15):m.end() + 12]
            out.append((bool(pre) and pre[0] in '전前', title, bool(RETRO.search(ctx)), ctx.replace('\n', ' '), m.start(), end, josa))
    return out



SPECIAL_SUN = re.compile(r'선거대책|선대위|선대본')
NONPOL = re.compile(r'(판사|검사|교수|변호사|작가|아나운서|앵커|기자|감독|선수|가수|배우|회장|이사장|원장|센터장|팀장|국장|처장|청장)$')


def clean(s):
    return re.sub(r'[\s,]+$', '', s.strip().strip('"').strip())


def parse_seed(text):
    """-> terms: {정규화 문자열: (대표호칭, 우선순위)}, jeon_listed: set(대표호칭), canon: [대표호칭 순서]"""
    rows = list(csv.reader(io.StringIO(text.lstrip('\ufeff'))))
    terms, canon, jeon_listed = {}, [], set()
    for r in rows[1:]:
        r = r + [''] * (3 - len(r))
        c0 = clean(r[0])
        if not c0:
            continue
        if c0.startswith('전 '):
            for x in re.split(r'\s*,\s*', c0):
                x = clean(x)
                if x.startswith('전'):
                    jeon_listed.add(x[1:].strip())
            continue
        if '선대위' in c0 or '선거대책위원회' in c0:
            canon.append('선대위 직책')
            continue
        canon.append(c0)
        terms[c0.replace(' ', '')] = (c0, 0)
    for r in rows[1:]:                                 # 변형은 대표호칭이 이미 있는 것에 밀린다
        r = r + [''] * (3 - len(r))
        c0 = clean(r[0])
        if not c0 or c0.startswith('전 ') or '선대위' in c0:
            continue
        for v in re.split(r'\s*,\s*', r[2]):
            v = re.sub(r'\([^)]*\)', '', v).strip()
            v = re.sub(r'O+(위원회)?\s*', '', v)        # OOO위원회 간사 -> 간사 (끝말로 맞춘다)
            v = v.replace(' ', '')
            if len(v) >= 2 and v not in terms:
                terms[v] = (c0, 1)
    return terms, canon, jeon_listed


SURNAMES = set('김이박최정강조윤장임한오서신권황안송류유홍전고문양손배백허남심노하곽성차주우구민나진지엄채천방공현함변염여추도소석선설마길연위표명기반왕금옥육인맹제모탁국어은편용예경봉사')
ORG_END = set('당부회청대처실국단도시군구')
ORG_STOP = {'감사원', '국정원', '대법원', '헌재', '검찰청'}
BAD_BEFORE = set('부내석동')      # 부대변인, 원내(대표), 수석(대변인), 공동(대표)처럼 다른 호칭을 만드는 앞글자


def looks_like_name(tok):
    return len(tok) == 3 and tok[0] in SURNAMES and tok[-1] not in ORG_END and tok not in ORG_STOP


def match_seed(title, terms):
    """호칭 문자열 -> (대표호칭 또는 None, 이름붙음 여부).
    끝말이 가장 길게 맞는 예시 호칭. 단, 앞에 다른 호칭이 붙은 복합 호칭(최고위원 후보, 부대변인)은 예시와 같은 호칭이 아니다."""
    parts = title.split()
    t = title.replace(' ', '')
    if SPECIAL_SUN.search(t):
        return '선대위 직책', False
    best = None
    for k, (c, pr) in terms.items():
        if t.endswith(k):
            pre = t[:-len(k)]
            if pre and pre[-1] in BAD_BEFORE:
                continue
            if pre and len(parts) >= 2:
                last_pre = parts[-2]
                if TITLE_SUFFIX.search(last_pre):
                    continue                      # 앞 낱말이 호칭(최고위원 후보)
                if looks_like_name(last_pre):
                    return None, True             # 윤건영 의원 — 옆 사람 이름
            key = (len(k), -pr)
            if best is None or key > best[0]:
                best = (key, c)
    if best is None and len(parts) >= 2 and looks_like_name(parts[-2]):
        return None, True
    return (best[1] if best else None), False


QOPEN, QCLOSE = '“', '”'


def sentence_spans(text):
    """(시작, 끝, 따옴표 안 구간 목록, 큰따옴표 포함 여부) 문장 목록.
    - 큰따옴표 안의 마침표는 문장 끝으로 보지 않는다.
    - 큰따옴표는 겹치지 않는다: 열린 상태에서 또 “ 가 나오면 닫는 따옴표의 오타(“…“)로 본다.
    - 줄바꿈에서는 열린 따옴표를 닫는다(여러 문단 인용문은 문단마다 다시 열린다).
    - 큰따옴표 밖에 있는 작은따옴표(‘…’) 구간도 따옴표 안으로 기록한다(남의 말·인용·제목)."""
    spans, st, qin, has_q = [], 0, [], False
    dopen = None        # 큰따옴표가 열린 위치
    sopen = None        # 작은따옴표가 열린 위치(큰따옴표 밖)
    n = len(text)
    i = 0
    while i < n:
        ch = text[i]
        if ch in '“"':
            if dopen is None:
                dopen = i; has_q = True; sopen = None
            else:                       # “…“ 오타, 또는 직선 큰따옴표 닫기
                qin.append((dopen, i)); dopen = None
        elif ch == '”':
            if dopen is not None:
                qin.append((dopen, i)); dopen = None
        elif ch == '‘' and dopen is None:
            sopen = i
        elif ch == '’' and dopen is None and sopen is not None:
            qin.append((sopen, i)); sopen = None
        elif ch == '\n':
            if dopen is not None:
                qin.append((dopen, i)); dopen = None
            sopen = None
            spans.append((st, i + 1, qin, has_q))
            st, qin, has_q = i + 1, [], False
        elif dopen is None and ch in '.?!' and (i + 1 >= n or text[i + 1] in ' \n”"\'’)'):
            spans.append((st, i + 1, qin, has_q))
            st, qin, has_q = i + 1, [], False
            sopen = None
        i += 1
    if st < n:
        if dopen is not None: qin.append((dopen, n))
        spans.append((st, n, qin, has_q))
    return spans


def in_quote_sentence(spans, pos):
    """pos가 속한 문장이 큰따옴표를 가지고 있고, pos 자신은 따옴표 안이 아닐 때 True."""
    for a, b, qin, has_q in spans:
        if a <= pos < b:
            if not has_q: return False
            return not any(x <= pos <= y for x, y in qin)
    return False


SUBJ = ('은', '는', '이', '가', '도', '만', '께서', '이라고', '라고')


def speaker_form(text, pos, tend, josa):
    """[이름 + 호칭]이 말하는 사람으로 쓰인 형태인가.
    - 호칭 뒤가 주격·보조사(은/는/이/가/도) 이거나, 조사 없이 바로 큰따옴표가 오고
    - 호칭 뒤 60자 안에 여는 따옴표가 있거나, 이름 앞 40자 안에 닫는 따옴표가 있어야 한다.
    '이준석 대표를 향해 “…”', '이준석 대표와 윤석열 후보의 갈등' 같은 남의 말·언급은 걸러진다."""
    if josa:
        if josa not in SUBJ:
            return False
    else:
        nxt = text[tend:tend + 3].lstrip()
        if not (nxt[:1] in ('“', '"', ':', ',', '(')):
            return False
    after = text[tend:tend + 60]
    before = text[max(0, pos - 40):pos]
    return any(c in after for c in '“"') or any(c in before for c in '”"')


def parity_inside(text, pos, tend):
    """편집인 방식: 같은 문단에서 이름+호칭 앞의 큰따옴표 수와 뒤의 큰따옴표 수가 둘 다 홀수면 인용문 안이다.
    “…“ 같은 오타가 있어도 개수는 변하지 않아서 따옴표의 모양에 흔들리지 않는다."""
    ps = text.rfind('\n', 0, pos) + 1
    pe = text.find('\n', tend)
    pe = len(text) if pe < 0 else pe
    b = sum(text[ps:pos].count(c) for c in '“”"')
    a = sum(text[tend:pe].count(c) for c in '“”"')
    return b % 2 == 1 and a % 2 == 1


def segments(items, gap=GAP):
    items = sorted(items)
    segs, cur = [], [items[0]]
    for x in items[1:]:
        if (x[0] - cur[-1][0]).days > gap:
            segs.append(cur); cur = [x]
        else:
            cur.append(x)
    segs.append(cur)
    return segs




_LOCK = threading.Lock()
BASE_SUFFIX = None


def _read_cols(header):
    di = header.index('날짜') if '날짜' in header else header.index('일자')
    ti = header.index('전문') if '전문' in header else header.index('본문')
    tt = header.index('제목')
    ui = header.index('URL') if 'URL' in header else None
    return di, ti, tt, ui


def analyze(header, data, name, seed_text, minart=3, adopted=(), isolated_gap=GAP):
    """header: 열 이름 목록, data: 행 목록. -> 결과 dict
    adopted: 2단계에서 사용자가 채택한 예시 밖 호칭(표기 그대로) 목록."""
    global TITLE_SUFFIX, BASE_SUFFIX
    with _LOCK:
        if BASE_SUFFIX is None:
            BASE_SUFFIX = TITLE_SUFFIX
        terms, canon, jeon_listed = parse_seed(seed_text)
        for a in adopted:                              # 채택한 호칭을 예시에 더한다
            a = a.strip()
            if not a: continue
            k = a.replace(' ', '')
            if k not in terms or terms[k][1] != 0:
                terms[k] = (a, 0)
            if a not in canon:
                canon.append(a)
        extra = sorted({k for k in terms if re.fullmatch(r'[가-힣]{2,8}', k)}, key=len, reverse=True)
        TITLE_SUFFIX = re.compile(r'(' + '|'.join(extra) + '|' + BASE_SUFFIX.pattern[1:-2] + r')$')
        try:
            return _analyze(header, data, name, terms, canon, minart, isolated_gap)
        finally:
            TITLE_SUFFIX = BASE_SUFFIX


def _analyze(header, data, name, terms, canon, minart, isolated_gap):
    di, ti, tt, ui = _read_cols(header)
    by_c = collections.defaultdict(lambda: {'cur': {}, 'jeon': {}, 'surf': collections.defaultdict(lambda: [0, 0])})
    other = collections.defaultdict(lambda: {'cur': {}, 'jeon': {}, 'ex': []})
    alld = []
    nonq = nq = 0
    exctx = {}
    excluded = []
    for r in data:
        d = norm_date(r[di])
        if not d or name not in r[ti] + r[tt]:
            continue
        alld.append(d)
        key = r[ui] if ui is not None else r[tt]
        txt = body_without_headings(r[ti])
        spans = sentence_spans(txt)
        for is_jeon, title, is_retro, ctx, pos, tend, josa in candidates(txt, name):
            if not (in_quote_sentence(spans, pos) and speaker_form(txt, pos, tend, josa)) or parity_inside(txt, pos, tend):
                nonq += 1
                continue
            nq += 1
            if re.search(r'당시\s*$', ctx.split(name)[0]) or title.startswith('당시 '):
                continue
            c, is_name = match_seed(title, terms)
            if is_name:
                continue
            side = 'jeon' if is_jeon else 'cur'
            if c:
                by_c[c][side][key] = (d, r[tt])
                exctx[(c, side, key)] = ctx
                by_c[c]['surf'][title][1 if is_jeon else 0] += 1
            else:
                o = other[title]
                o[side][key] = (d, r[tt])
                if len(o['ex']) < 2:
                    o['ex'].append(ctx)
    if not alld:
        return {'error': f'기사에서 "{name}"을(를) 찾지 못했습니다.'}
    first, last = min(alld), max(alld)

    T1 = ['호칭', '구분', '구간', '시작일', '종료일', '구간 기사수', '이 호칭 전체 기사수',
          '구간 끝 이후 첫 "전 호칭" 등장일', '표기 변형(건수)', '비고', '시작 근거 기사', '종료 근거 기사']
    out_rows = []
    for c in canon:
        dat = by_c.get(c)
        if not dat or (not dat['cur'] and not dat['jeon']):
            continue
        for side, label in (('cur', '현직'), ('jeon', '전(前)')):
            arts = dat[side]
            if len(arts) < minart:
                continue
            jdates = sorted(d for d, _ in dat['jeon'].values())
            segs = segments(arts.values(), isolated_gap if side == 'cur' else isolated_gap * 3)
            alld_dates = sorted(d for d, _ in arts.values())
            keep = []
            for sg in segs:
                if len(sg) == 1:
                    d1, tit1 = sg[0]
                    k1 = next(k for k, v in arts.items() if v == (d1, tit1))
                    near = min(abs((d1 - x).days) for x in alld_dates if x != d1) if len(set(alld_dates)) > 1 else ''
                    excluded.append([c if side == 'cur' else '전 ' + c, label, d1.isoformat(), near, tit1[:50], exctx.get((c, side, k1), '')])
                else:
                    keep.append(sg)
            for k, sg in enumerate(keep, 1):
                a, b = sg[0][0], sg[-1][0]
                note = []
                if c == '최고위원' and side == 'cur': note.append('수석최고위원도 "최고위원"으로 불리므로 수석 기간을 포함')
                if a == first: note.append('자료 시작 전부터일 수 있음')
                if b == last: note.append('자료 끝 이후에도 이어질 수 있음')
                if side == 'cur':
                    jin = sum(1 for d in jdates if a <= d <= b)
                    if jin >= 3 and jin > len(sg):
                        note.append(f'구간 안 "전 {c}" {jin}건이 현직형 {len(sg)}건보다 많음 — 구간 전체가 현직 시기가 아닐 수 있음')
                    after = next((d for d in jdates if d > b), None)
                else:
                    after = None
                surf = '; '.join(f'{s}({v[1] if side == "jeon" else v[0]})' for s, v in
                                 sorted(dat['surf'].items(), key=lambda kv: -(kv[1][1] if side == 'jeon' else kv[1][0]))
                                 if (v[1] if side == 'jeon' else v[0]))
                out_rows.append((a, canon.index(c), [c if side == 'cur' else '전 ' + c, label, k, a.isoformat(), b.isoformat(),
                                                    len(sg), len(arts), after.isoformat() if after else '', surf, ', '.join(note),
                                                    f'{sg[0][0]} {sg[0][1][:40]}', f'{sg[-1][0]} {sg[-1][1][:40]}']))
    t1 = [r for _a, _i, r in sorted(out_rows, key=lambda x: (x[0], x[1]))]
    T2 = ['선정(O/X)', '호칭', '현직형 기사수', '현직형 첫 등장', '현직형 마지막 등장', '"전 호칭" 기사수', '"전" 첫 등장', '"전" 마지막 등장', '주의', '예문1', '예문2']
    lst = []
    for t, o in other.items():
        n, nj = len(o['cur']), len(o['jeon'])
        if n < minart and nj < minart:
            continue
        cd = sorted(d for d, _ in o['cur'].values())
        jd = sorted(d for d, _ in o['jeon'].values())
        warn = []
        if NONPOL.search(t.replace(' ', '')): warn.append('비정치 직종 — 동명이인 또는 경력 소개 가능성')
        if n >= minart and n < 5 and not nj: warn.append('기사 적음')
        lst.append((-(n + nj), ['', t, n, cd[0].isoformat() if cd else '', cd[-1].isoformat() if cd else '',
                                nj, jd[0].isoformat() if jd else '', jd[-1].isoformat() if jd else '', ', '.join(warn)] + o['ex'] + [''] * (2 - len(o['ex']))))
    t2 = [r for _k, r in sorted(lst, key=lambda x: x[0])]
    T3 = ['호칭', '구분', '날짜', '가장 가까운 다른 날짜와의 거리(일)', '기사 제목', '앞뒤 문맥']
    return {'name': name, 'T1': T1, 't1': t1, 'T2': T2, 't2': t2, 'T3': T3, 't3': sorted(excluded, key=lambda x: (x[0], x[2])),
            'stats': {'mentions': len(alld), 'caught': nq + nonq, 'used': nq, 'first': first.isoformat(), 'last': last.isoformat()}}


def build_xlsx(res):
    """결과 dict -> 엑셀 파일 바이트"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    sheets = [('1.호칭 기간', res['T1'], res['t1'], (22, 8, 6, 12, 12, 10, 12, 18, 40, 40, 44, 44)),
              ('2.예시에 없는 호칭(선정)', res['T2'], res['t2'], (10, 26, 10, 14, 14, 12, 14, 14, 40, 60, 60)),
              ('배제된 외톨이 1건', res['T3'], res['t3'], (22, 8, 12, 14, 40, 70))]
    for i, (title, hdr, rows, widths) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(hdr)
        for r in rows: ws.append(r)
        for c in ws[1]: c.font = Font(bold=True)
        ws.freeze_panes = 'A2'
        for j, w in enumerate(widths):
            ws.column_dimensions[chr(65 + j)].width = w
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def read_table(file_bytes, filename):
    """업로드된 csv/xlsx -> (header, rows)"""
    if filename.lower().endswith(('.xlsx', '.xls')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        rows = [[('' if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        for enc in ('utf-8-sig', 'cp949'):
            try:
                text = file_bytes.decode(enc); break
            except UnicodeDecodeError:
                continue
        csv.field_size_limit(10**9)
        rows = list(csv.reader(io.StringIO(text)))
    return rows[0], [r for r in rows[1:] if any(r)]


# ───────────────────────── 동성 동호칭 타인 기사 찾기 ─────────────────────────
# 정의(편집인 확정, 2026-10-18): 지정발언자와 성이 같고 이름이 다른 사람이 '정치 호칭'과 함께 같은 기사에 나오면
# 그 기사는 '풀네임 확인 필요' 기사다. 정치인은 겸직이 흔하므로 호칭은 정확히 같지 않아도 겸직 계열이면 같은 호칭으로 본다.
FAMILY_DEFAULT = ('의원, 최고위원, 위원, 위원장, 간사, 대표, 당대표, 원내대표, 부대표, 원내부대표, 수석부대표, '
                  '원내수석부대표, 원내정책수석부대표, 정책수석부대표, 원내수석, 원내정책수석, 정책위의장, 사무총장, '
                  '대변인, 수석대변인, 부대변인, 원내대변인, 의장, 부의장, 비대위원장, 혁신위원장, 후보')
_PARTY = (r'(?:(?:더불어민주당|민주당|국민의힘|국힘|조국혁신당|혁신당|개혁신당|진보당|기본소득당|사회민주당|'
          r'새로운미래|무소속|야당|여당|당)\s)')
SAME_COLS = ['일자', '제목', '타인 이름', '호칭', '기사에 지정발언자 포함', '문장', 'URL']
SAME_SUM_COLS = ['타인 이름', '기사수', '호칭', '첫 등장', '마지막 등장', '지정발언자와 함께 나온 기사']


def find_same_surname(header, data, name, titles_text=FAMILY_DEFAULT):
    di, ti, tt, ui = _read_cols(header)
    titles = [t.strip() for t in re.split(r'[,\n]', titles_text) if t.strip()]
    if not titles or not name.strip():
        return {'error': '호칭 목록과 성명을 입력하세요.'}
    name = name.strip()
    T = '|'.join(sorted(map(re.escape, titles), key=len, reverse=True))
    pat = re.compile(r'(?<![가-힣])(' + re.escape(name[0]) + r'[가-힣]{2})\s?' + _PARTY + r'{0,2}(전\s)?(' + T + r')(?![가-힣])')
    rows, seen = [], set()
    for r in data:
        if len(r) <= max(di, ti, tt):
            continue
        text = body_without_headings(r[ti])
        hits = {}
        for m in pat.finditer(text):
            nm, jeon, ti_ = m.group(1), m.group(2), m.group(3)
            if nm == name or jeon or not looks_like_name(nm):
                continue
            a = max(0, m.start() - 30)
            hits.setdefault((nm, ti_), text[a:m.end() + 40].replace('\n', ' '))
        if not hits:
            continue
        url = r[ui] if ui is not None and ui < len(r) else ''
        key = url or (r[di], r[tt])
        if key in seen:
            continue
        seen.add(key)
        has = 'O' if name in text else ''
        for (nm, ti_), ctx in hits.items():
            rows.append([norm_date(r[di]), r[tt], nm, ti_, has, ctx, url])
    rows.sort(key=lambda x: (x[0], x[1]))
    by = {}
    for x in rows:
        b = by.setdefault(x[2], {'urls': set(), 'ti': set(), 'dates': [], 'with': 0})
        b['urls'].add(x[6] or x[1]); b['ti'].add(x[3]); b['dates'].append(x[0]); b['with'] += (x[4] == 'O')
    summ = sorted(([k, len(v['urls']), ', '.join(sorted(v['ti'])), min(v['dates']), max(v['dates']), v['with']]
                   for k, v in by.items()), key=lambda s: -s[1])
    return {'T': SAME_COLS, 'rows': rows, 'TS': SAME_SUM_COLS, 'summ': summ, 'n_articles': len(seen), 'total': len(data)}


def same_xlsx(res):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active; ws.title = '1.기사 목록'
    ws.append(res['T']); [ws.append(r) for r in res['rows']]
    ws2 = wb.create_sheet('2.타인별 요약')
    ws2.append(res['TS']); [ws2.append(r) for r in res['summ']]
    for w, widths in ((ws, (11, 40, 10, 14, 10, 60, 40)), (ws2, (10, 8, 24, 11, 11, 12))):
        for i, wd in enumerate(widths):
            w.column_dimensions[chr(65 + i)].width = wd
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()


def same_csv(res):
    """매직박스에 올리는 용도: URL이 있는 행만. 열: URL, 일자, 제목, 타인 이름, 호칭"""
    buf = io.StringIO(); w = csv.writer(buf)
    w.writerow(['URL', '일자', '제목', '타인 이름', '호칭'])
    seen = set()
    for r in res['rows']:
        if r[6] and r[6] not in seen:
            seen.add(r[6]); w.writerow([r[6], r[0], r[1], r[2], r[3]])
    return buf.getvalue().encode('utf-8-sig')
